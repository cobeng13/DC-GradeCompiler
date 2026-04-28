import argparse
import csv
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, TextIO, Tuple

from openpyxl import Workbook, load_workbook


PASSED_COLUMNS = [
    "SourceFile",
    "Section",
    "Course",
    "StudentNumber",
    "StudentName",
    "MidtermExam",
    "MidtermGrade",
    "Status",
    "Notes",
]

ISSUE_COLUMNS = [
    "SourceFile",
    "RowNumber",
    "IssueType",
    "Details",
    "RawStudentNumber",
    "RawStudentName",
    "RawMidtermExam",
    "RawMidtermGrade",
]

HEADER_ALIASES = {
    "StudentNumber": [
        "StudentNumber",
        "Student Number",
        "Student No",
        "Student No.",
        "Student ID",
    ],
    "StudentName": [
        "StudentName",
        "Student Name",
        "Name",
    ],
    "MidtermExam": [
        "MidtermExam",
        "Midterm Exam",
        "ME",
        "Midterm Score",
    ],
    "MidtermGrade": [
        "MidtermGrade",
        "Midterm Grade",
        "MT Grade",
        "Midterm",
    ],
}

MASTERLIST_ID_ALIASES = [
    "StudentNumber",
    "Student Number",
    "Student No",
    "Student No.",
    "Student ID",
    "ClassNumber",
    "Class Number",
]


def parse_filename(filename: str) -> Tuple[str, str, Optional[str]]:
    """Parse Section and Course from names like 2A_OrgMedChem.xlsm."""
    path = Path(filename)
    stem = path.stem

    if "_" not in stem:
        return "", "", "Filename does not follow Section_Course pattern."

    section, course = stem.split("_", 1)
    if not section.strip() or not course.strip():
        return "", "", "Filename does not follow Section_Course pattern."

    return section.strip(), course.strip(), None


def normalize_header(value: Any) -> str:
    """Normalize a header for alias matching."""
    if value is None:
        return ""

    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", "", text)
    return text


def normalize_name(value: Any) -> str:
    """Normalize names for exact matching across punctuation and spacing changes."""
    if value is None:
        return ""

    text = str(value).strip().lower()
    tokens = re.findall(r"[a-z0-9]+", text)
    return " ".join(tokens)


def _alias_lookup() -> Dict[str, str]:
    lookup = {}
    for canonical_name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            lookup[normalize_header(alias)] = canonical_name
    return lookup


def find_column_map(header_row: Tuple[Any, ...]) -> Dict[str, int]:
    """Return a map of canonical header names to one-based column indexes."""
    lookup = _alias_lookup()
    column_map: Dict[str, int] = {}

    for index, cell_value in enumerate(header_row, start=1):
        normalized = normalize_header(cell_value)
        canonical_name = lookup.get(normalized)
        if canonical_name and canonical_name not in column_map:
            column_map[canonical_name] = index

    return column_map


def is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def row_is_blank(values: List[Any]) -> bool:
    return all(is_blank(value) for value in values)


def student_row_is_blank(
    student_number: Any, student_name: Any, midterm_exam: Any
) -> bool:
    """Skip template/formula rows that do not identify a real student."""
    return is_blank(student_number) and is_blank(student_name) and is_blank(midterm_exam)


def to_number(value: Any) -> Optional[float]:
    """Convert grades that may be stored as numbers or text."""
    if is_blank(value):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1].strip()

    try:
        return float(text)
    except ValueError:
        return None


def get_cell_value(row: Tuple[Any, ...], column_map: Dict[str, int], key: str) -> Any:
    column_index = column_map.get(key)
    if column_index is None or column_index > len(row):
        return None
    return row[column_index - 1]


def _read_csv_sample(file_obj: TextIO) -> str:
    sample = file_obj.read(4096)
    file_obj.seek(0)
    return sample


def find_header_by_alias(header_map: Dict[str, str], aliases: List[str]) -> Optional[str]:
    for alias in aliases:
        header = header_map.get(normalize_header(alias))
        if header:
            return header
    return None


def load_masterlist(masterlist_path: Path) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    """Load masterlist.csv into a normalized-name to student-number lookup."""
    issues: List[Dict[str, Any]] = []
    masterlist: Dict[str, Any] = {}

    if not masterlist_path.exists():
        issues.append(
            issue_row(
                masterlist_path.name,
                "",
                "MasterlistMissing",
                f"Masterlist file does not exist: {masterlist_path}",
            )
        )
        return masterlist, issues

    try:
        with masterlist_path.open("r", newline="", encoding="utf-8-sig") as file_obj:
            sample = _read_csv_sample(file_obj)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t")
            except csv.Error:
                dialect = csv.excel

            reader = csv.DictReader(file_obj, dialect=dialect)
            if reader.fieldnames is None:
                issues.append(
                    issue_row(
                        masterlist_path.name,
                        1,
                        "MasterlistInvalid",
                        "Masterlist has no header row.",
                    )
                )
                return masterlist, issues

            header_map = {
                normalize_header(header): header for header in reader.fieldnames if header
            }
            name_column = header_map.get(normalize_header("Name"))
            student_number_column = find_header_by_alias(
                header_map, MASTERLIST_ID_ALIASES
            )

            if not name_column or not student_number_column:
                issues.append(
                    issue_row(
                        masterlist_path.name,
                        1,
                        "MasterlistInvalid",
                        "Masterlist must contain Name and an ID column such as StudentNumber or ClassNumber.",
                    )
                )
                return masterlist, issues

            for row_number, row in enumerate(reader, start=2):
                raw_name = row.get(name_column)
                raw_student_number = row.get(student_number_column)
                normalized_name = normalize_name(raw_name)

                if not normalized_name:
                    continue

                if normalized_name in masterlist:
                    if is_blank(masterlist[normalized_name]) and not is_blank(
                        raw_student_number
                    ):
                        masterlist[normalized_name] = raw_student_number
                    continue

                masterlist[normalized_name] = raw_student_number
    except Exception as exc:
        issues.append(
            issue_row(
                masterlist_path.name,
                "",
                "MasterlistReadError",
                f"Could not read masterlist: {exc}",
            )
        )

    return masterlist, issues


def find_masterlist_student_number(
    masterlist: Dict[str, Any], student_name: Any
) -> Optional[Any]:
    normalized_name = normalize_name(student_name)
    if not normalized_name:
        return None
    return masterlist.get(normalized_name)


def issue_row(
    source_file: str,
    row_number: Any,
    issue_type: str,
    details: str,
    raw_student_number: Any = None,
    raw_student_name: Any = None,
    raw_midterm_exam: Any = None,
    raw_midterm_grade: Any = None,
) -> Dict[str, Any]:
    return {
        "SourceFile": source_file,
        "RowNumber": row_number,
        "IssueType": issue_type,
        "Details": details,
        "RawStudentNumber": raw_student_number,
        "RawStudentName": raw_student_name,
        "RawMidtermExam": raw_midterm_exam,
        "RawMidtermGrade": raw_midterm_grade,
    }


def process_workbook(
    workbook_path: Path, passing_grade: float, masterlist: Dict[str, Any]
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], bool]:
    """
    Read a grading workbook and return passed student rows, issue rows, and
    whether at least one worksheet was processable.
    """
    passed_students: List[Dict[str, Any]] = []
    issues: List[Dict[str, Any]] = []
    source_file = workbook_path.name
    section, course, filename_note = parse_filename(source_file)

    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
            keep_vba=workbook_path.suffix.lower() == ".xlsm",
        )
    except Exception as exc:
        issues.append(
            issue_row(
                source_file,
                "",
                "WorkbookReadError",
                f"Could not open workbook: {exc}",
            )
        )
        return passed_students, issues, False

    processed_any_sheet = False

    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None or row_is_blank(list(header_row)):
            continue

        column_map = find_column_map(header_row)
        hard_required = ["StudentName", "MidtermExam", "MidtermGrade"]
        missing_required = [name for name in hard_required if name not in column_map]

        if missing_required:
            issues.append(
                issue_row(
                    source_file,
                    1,
                    "MissingRequiredColumns",
                    f"Sheet '{worksheet.title}' is missing: {', '.join(missing_required)}",
                )
            )
            continue

        processed_any_sheet = True

        if "StudentNumber" not in column_map:
            issues.append(
                issue_row(
                    source_file,
                    1,
                    "MissingStudentNumberColumn",
                    f"Sheet '{worksheet.title}' has no StudentNumber column; rows will be flagged.",
                )
            )

        for row_number, row in enumerate(rows, start=2):
            row_values = list(row)
            if row_is_blank(row_values):
                continue

            raw_student_number = get_cell_value(row, column_map, "StudentNumber")
            raw_student_name = get_cell_value(row, column_map, "StudentName")
            raw_midterm_exam = get_cell_value(row, column_map, "MidtermExam")
            raw_midterm_grade = get_cell_value(row, column_map, "MidtermGrade")

            if student_row_is_blank(
                raw_student_number, raw_student_name, raw_midterm_exam
            ):
                continue

            grade = to_number(raw_midterm_grade)
            if grade is None:
                issues.append(
                    issue_row(
                        source_file,
                        row_number,
                        "InvalidMidtermGrade",
                        f"Sheet '{worksheet.title}' has a blank or non-numeric MidtermGrade.",
                        raw_student_number,
                        raw_student_name,
                        raw_midterm_exam,
                        raw_midterm_grade,
                    )
                )
                continue

            if grade < passing_grade:
                continue

            notes = []
            if filename_note:
                notes.append(filename_note)

            student_number = raw_student_number
            if "StudentNumber" not in column_map:
                masterlist_student_number = find_masterlist_student_number(
                    masterlist, raw_student_name
                )
                if not is_blank(masterlist_student_number):
                    student_number = masterlist_student_number
                    notes.append("StudentNumber column missing; filled from masterlist.")
                else:
                    notes.append("StudentNumber column missing; no masterlist match.")
            elif is_blank(raw_student_number):
                masterlist_student_number = find_masterlist_student_number(
                    masterlist, raw_student_name
                )
                if not is_blank(masterlist_student_number):
                    student_number = masterlist_student_number
                    notes.append("StudentNumber filled from masterlist.")
                else:
                    notes.append("StudentNumber is blank; no masterlist match.")

            passed_students.append(
                {
                    "SourceFile": source_file,
                    "Section": section,
                    "Course": course,
                    "StudentNumber": student_number,
                    "StudentName": raw_student_name,
                    "MidtermExam": raw_midterm_exam,
                    "MidtermGrade": grade,
                    "Status": "Passed",
                    "Notes": " ".join(notes),
                }
            )

    workbook.close()

    if not processed_any_sheet:
        issues.append(
            issue_row(
                source_file,
                "",
                "NoProcessableSheet",
                "No worksheet contained all required processable columns.",
            )
        )

    return passed_students, issues, processed_any_sheet


def write_sheet(workbook: Workbook, title: str, columns: List[str], rows: List[Dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title=title)
    worksheet.append(columns)

    for row in rows:
        worksheet.append([row.get(column, "") for column in columns])

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 12), 60
        )


def generate_report(
    input_folder: Path,
    output_file: Path,
    passing_grade: float,
    masterlist_path: Path,
) -> Tuple[int, int, int]:
    """Process all grading sheets and write the final report workbook."""
    passed_students: List[Dict[str, Any]] = []
    issues: List[Dict[str, Any]] = []
    files_processed = 0

    output_file.parent.mkdir(parents=True, exist_ok=True)
    masterlist, masterlist_issues = load_masterlist(masterlist_path)
    issues.extend(masterlist_issues)

    if not input_folder.exists():
        issues.append(
            issue_row(
                "",
                "",
                "InputFolderMissing",
                f"Input folder does not exist: {input_folder}",
            )
        )
    else:
        workbook_paths = sorted(
            path
            for path in input_folder.iterdir()
            if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm"}
        )

        if not workbook_paths:
            issues.append(
                issue_row(
                    "",
                    "",
                    "NoInputFiles",
                    f"No .xlsx or .xlsm files found in: {input_folder}",
                )
            )

        for workbook_path in workbook_paths:
            workbook_passed, workbook_issues, processed = process_workbook(
                workbook_path, passing_grade, masterlist
            )
            if processed:
                files_processed += 1
            passed_students.extend(workbook_passed)
            issues.extend(workbook_issues)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    write_sheet(workbook, "PassedStudents", PASSED_COLUMNS, passed_students)
    write_sheet(workbook, "Issues", ISSUE_COLUMNS, issues)
    workbook.save(output_file)

    return files_processed, len(passed_students), len(issues)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate a report of students who passed the midterm exam."
    )
    parser.add_argument(
        "--input",
        default="input/grades",
        help="Folder containing .xlsx and .xlsm grading sheets.",
    )
    parser.add_argument(
        "--output",
        default="output/passed_midterm_report.xlsx",
        help="Path for the generated Excel report.",
    )
    parser.add_argument(
        "--passing-grade",
        type=float,
        default=75,
        help="Minimum MidtermGrade required to pass.",
    )
    parser.add_argument(
        "--masterlist",
        default="input/masterlist/masterlist.csv",
        help="Path to masterlist.csv used to fill missing student numbers.",
    )
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    input_folder = Path(args.input)
    output_file = Path(args.output)
    masterlist_path = Path(args.masterlist)
    files_processed, passed_count, issue_count = generate_report(
        input_folder, output_file, args.passing_grade, masterlist_path
    )

    print("Midterm report generated.")
    print(f"Files processed: {files_processed}")
    print(f"Passed students found: {passed_count}")
    print(f"Issues found: {issue_count}")
    print(f"Output file: {output_file.resolve()}")


if __name__ == "__main__":
    main()
