import csv
import json
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, TextIO, Tuple

from openpyxl import load_workbook


PASSED_COLUMNS = [
    "SourceFile",
    "Section",
    "Course",
    "StudentNumber",
    "StudentName",
    "MidtermExam",
    "MidtermGrade",
    "Status",
    "Notes",
]

STANDARD_GRADE_COLUMNS = [
    "ClassNumber",
    "MQ1",
    "MQ2",
    "MQ3",
    "MQ4",
    "MQ5",
    "MQ6",
    "MQ7",
    "MQ8",
    "MQ9",
    "MQ10",
    "MQAve",
    "MQAve70",
    "MA1",
    "MA2",
    "MA3",
    "MA4",
    "MA5",
    "MA6",
    "MA7",
    "MA8",
    "MA9",
    "MA10",
    "MActAve",
    "MActAve25",
    "MResProj",
    "MResProj5",
    "MClassPerf",
    "MClassPerf30",
    "PrelimExam",
    "30PrelimExam",
    "MidtermExam",
    "MidtermExam40",
    "MidtermGrade",
    "MidtermGrade40",
    "FQ1",
    "FQ2",
    "FQ3",
    "FQ4",
    "FQ5",
    "FQ6",
    "FQ7",
    "FQ8",
    "FQ9",
    "FQ10",
    "FQAve",
    "FQAve70",
    "FA1",
    "FA2",
    "FA3",
    "FA4",
    "FA5",
    "FA6",
    "FA7",
    "FA8",
    "FA9",
    "FA10",
    "FActAve",
    "FActAve25",
    "FResProj",
    "FResProj5",
    "FClassPerf",
    "FClassPerf20",
    "FinalExam",
    "FinalExam40",
    "FinalGrade",
    "Interpretation",
]

ISSUE_COLUMNS = [
    "SourceFile",
    "RowNumber",
    "IssueType",
    "Details",
    "RawStudentNumber",
    "RawStudentName",
    "RawMidtermExam",
    "RawMidtermGrade",
]

HEADER_ALIASES = {
    "StudentNumber": [
        "StudentNumber",
        "Student Number",
        "Student No",
        "Student No.",
        "Student ID",
    ],
    "StudentName": [
        "StudentName",
        "Student Name",
        "Name",
    ],
    "MidtermExam": [
        "MidtermExam",
        "Midterm Exam",
        "ME",
        "Midterm Score",
    ],
    "MidtermGrade": [
        "MidtermGrade",
        "Midterm Grade",
        "MT Grade",
        "Midterm",
    ],
}

for standard_column in STANDARD_GRADE_COLUMNS:
    HEADER_ALIASES.setdefault(standard_column, []).append(standard_column)

MASTERLIST_ID_ALIASES = [
    "StudentNumber",
    "Student Number",
    "Student No",
    "Student No.",
    "Student ID",
    "ClassNumber",
    "Class Number",
]

MASTERLIST_SECTION_ALIASES = [
    "Section",
    "Class",
    "Class Section",
]

TEXT_REPLACEMENTS = {
    "\u5e3d": "n",  # Observed bad substitution for enye in faculty/masterlist exports.
    "Ã‘": "Ñ",
    "Ã’": "Ñ",
    "Ã\u2018": "Ñ",
    "Ã\u2019": "Ñ",
    "Ã±": "ñ",
}


@dataclass
class IngestResult:
    passing_grade: float
    source_files: List[Dict[str, Any]] = field(default_factory=list)
    result_rows: List[Dict[str, Any]] = field(default_factory=list)
    issues: List[Dict[str, Any]] = field(default_factory=list)
    masterlist_records: List[Dict[str, Any]] = field(default_factory=list)
    files_processed: int = 0

    @property
    def passed_count(self) -> int:
        return len([row for row in self.result_rows if row.get("Status") == "Passed"])


def parse_filename(filename: str) -> Tuple[str, str, Optional[str]]:
    path = Path(filename)
    stem = path.stem

    if "_" not in stem:
        return "", "", "Filename does not follow Section_Course pattern."

    section, course = stem.split("_", 1)
    if not section.strip() or not course.strip():
        return "", "", "Filename does not follow Section_Course pattern."

    return section.strip(), course.strip(), None


def normalize_header(value: Any) -> str:
    if value is None:
        return ""

    text = strip_diacritics(repair_text_encoding(str(value))).strip().lower()
    return re.sub(r"[^a-z0-9]+", "", text)


def repair_text_encoding(value: str) -> str:
    text = value
    for bad_text, replacement in TEXT_REPLACEMENTS.items():
        text = text.replace(bad_text, replacement)

    if "\u00c3" in text:
        try:
            repaired = text.encode("latin-1").decode("utf-8")
        except UnicodeError:
            return text
        return repaired

    return text


def strip_diacritics(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(
        character
        for character in normalized
        if not unicodedata.combining(character)
    )


def clean_student_name(value: Any) -> str:
    if value is None:
        return ""

    text = repair_text_encoding(str(value)).strip()
    numbered_name = re.match(r"^\s*\d+\s+(.+)$", text)
    if numbered_name and "," in numbered_name.group(1):
        text = numbered_name.group(1).strip()

    text = re.sub(r"\d+", "", text)
    return re.sub(r"\s+", " ", text)


def normalize_name(value: Any) -> str:
    if value is None:
        return ""

    text = strip_diacritics(clean_student_name(value)).strip().lower()
    tokens = re.findall(r"[a-z]+", text)
    return " ".join(tokens)


def _alias_lookup() -> Dict[str, str]:
    lookup = {}
    for canonical_name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            lookup[normalize_header(alias)] = canonical_name
    return lookup


def find_column_map(header_row: Tuple[Any, ...]) -> Dict[str, int]:
    lookup = _alias_lookup()
    column_map: Dict[str, int] = {}

    for index, cell_value in enumerate(header_row, start=1):
        canonical_name = lookup.get(normalize_header(cell_value))
        if canonical_name and canonical_name not in column_map:
            column_map[canonical_name] = index

    return column_map


def is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def row_is_blank(values: List[Any]) -> bool:
    return all(is_blank(value) for value in values)


def student_row_is_blank(
    student_number: Any, student_name: Any, midterm_exam: Any
) -> bool:
    return is_blank(student_number) and is_blank(student_name) and is_blank(midterm_exam)


def to_number(value: Any) -> Optional[float]:
    if is_blank(value):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1].strip()

    try:
        return float(text)
    except ValueError:
        return None


def get_cell_value(row: Tuple[Any, ...], column_map: Dict[str, int], key: str) -> Any:
    column_index = column_map.get(key)
    if column_index is None or column_index > len(row):
        return None
    return row[column_index - 1]


def _read_csv_sample(file_obj: TextIO) -> str:
    sample = file_obj.read(4096)
    file_obj.seek(0)
    return sample


def find_header_by_alias(header_map: Dict[str, str], aliases: List[str]) -> Optional[str]:
    for alias in aliases:
        header = header_map.get(normalize_header(alias))
        if header:
            return header
    return None


def issue_row(
    source_file: str,
    row_number: Any,
    issue_type: str,
    details: str,
    raw_student_number: Any = None,
    raw_student_name: Any = None,
    raw_midterm_exam: Any = None,
    raw_midterm_grade: Any = None,
) -> Dict[str, Any]:
    return {
        "SourceFile": source_file,
        "RowNumber": row_number,
        "IssueType": issue_type,
        "Details": details,
        "RawStudentNumber": raw_student_number,
        "RawStudentName": raw_student_name,
        "RawMidtermExam": raw_midterm_exam,
        "RawMidtermGrade": raw_midterm_grade,
    }


def load_masterlist(
    masterlist_path: Path,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]], List[Dict[str, Any]]]:
    issues: List[Dict[str, Any]] = []
    masterlist: Dict[str, Any] = {}
    masterlist_records: List[Dict[str, Any]] = []
    record_indexes: Dict[str, int] = {}

    if not masterlist_path.exists():
        issues.append(
            issue_row(
                masterlist_path.name,
                "",
                "MasterlistMissing",
                f"Masterlist file does not exist: {masterlist_path}",
            )
        )
        return masterlist, masterlist_records, issues

    try:
        with masterlist_path.open("r", newline="", encoding="utf-8-sig") as file_obj:
            sample = _read_csv_sample(file_obj)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t")
            except csv.Error:
                dialect = csv.excel

            reader = csv.DictReader(file_obj, dialect=dialect)
            if reader.fieldnames is None:
                issues.append(
                    issue_row(
                        masterlist_path.name,
                        1,
                        "MasterlistInvalid",
                        "Masterlist has no header row.",
                    )
                )
                return masterlist, masterlist_records, issues

            header_map = {
                normalize_header(header): header for header in reader.fieldnames if header
            }
            name_column = header_map.get(normalize_header("Name"))
            student_number_column = find_header_by_alias(
                header_map, MASTERLIST_ID_ALIASES
            )
            section_column = find_header_by_alias(header_map, MASTERLIST_SECTION_ALIASES)

            if not name_column or not student_number_column:
                issues.append(
                    issue_row(
                        masterlist_path.name,
                        1,
                        "MasterlistInvalid",
                        "Masterlist must contain Name and an ID column such as StudentNumber or ClassNumber.",
                    )
                )
                return masterlist, masterlist_records, issues

            for row in reader:
                raw_name = row.get(name_column)
                raw_student_number = row.get(student_number_column)
                raw_section = row.get(section_column) if section_column else ""
                normalized_name = normalize_name(raw_name)

                if not normalized_name:
                    continue

                if normalized_name in masterlist:
                    record = masterlist_records[record_indexes[normalized_name]]
                    if is_blank(masterlist[normalized_name]) and not is_blank(
                        raw_student_number
                    ):
                        masterlist[normalized_name] = raw_student_number
                        record["StudentNumber"] = raw_student_number
                    if is_blank(record.get("Section")) and not is_blank(raw_section):
                        record["Section"] = raw_section
                    continue

                masterlist[normalized_name] = raw_student_number
                record_indexes[normalized_name] = len(masterlist_records)
                masterlist_records.append(
                    {
                        "StudentName": raw_name,
                        "StudentNumber": raw_student_number,
                        "Section": raw_section,
                        "NormalizedName": normalized_name,
                    }
                )
    except Exception as exc:
        issues.append(
            issue_row(
                masterlist_path.name,
                "",
                "MasterlistReadError",
                f"Could not read masterlist: {exc}",
            )
        )

    return masterlist, masterlist_records, issues


def find_masterlist_student_number(
    masterlist: Dict[str, Any], student_name: Any
) -> Optional[Any]:
    normalized_name = normalize_name(student_name)
    if not normalized_name:
        return None
    return masterlist.get(normalized_name)


def process_workbook(
    workbook_path: Path, passing_grade: float, masterlist: Dict[str, Any]
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Dict[str, Any]]:
    result_rows: List[Dict[str, Any]] = []
    issues: List[Dict[str, Any]] = []
    source_file = workbook_path.name
    section, course, filename_note = parse_filename(source_file)
    source_info = {
        "SourceFile": source_file,
        "Path": str(workbook_path),
        "Section": section,
        "Course": course,
        "Processed": False,
        "Notes": filename_note or "",
    }

    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
            keep_vba=workbook_path.suffix.lower() == ".xlsm",
        )
    except Exception as exc:
        issues.append(
            issue_row(
                source_file,
                "",
                "WorkbookReadError",
                f"Could not open workbook: {exc}",
            )
        )
        return result_rows, issues, source_info

    processed_any_sheet = False

    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None or row_is_blank(list(header_row)):
            continue

        column_map = find_column_map(header_row)
        hard_required = ["StudentName", "MidtermExam", "MidtermGrade"]
        missing_required = [name for name in hard_required if name not in column_map]

        if missing_required:
            issues.append(
                issue_row(
                    source_file,
                    1,
                    "MissingRequiredColumns",
                    f"Sheet '{worksheet.title}' is missing: {', '.join(missing_required)}",
                )
            )
            continue

        processed_any_sheet = True

        if "StudentNumber" not in column_map:
            issues.append(
                issue_row(
                    source_file,
                    1,
                    "MissingStudentNumberColumn",
                    f"Sheet '{worksheet.title}' has no StudentNumber column; rows will be flagged.",
                )
            )

        for row_number, row in enumerate(rows, start=2):
            row_values = list(row)
            if row_is_blank(row_values):
                continue

            raw_student_number = get_cell_value(row, column_map, "StudentNumber")
            raw_student_name = get_cell_value(row, column_map, "StudentName")
            student_name = clean_student_name(raw_student_name)
            raw_midterm_exam = get_cell_value(row, column_map, "MidtermExam")
            raw_midterm_grade = get_cell_value(row, column_map, "MidtermGrade")
            standard_grade_values = {
                column: get_cell_value(row, column_map, column)
                for column in STANDARD_GRADE_COLUMNS
            }

            if student_row_is_blank(
                raw_student_number, raw_student_name, raw_midterm_exam
            ):
                continue

            grade = to_number(raw_midterm_grade)
            if grade is None:
                issues.append(
                    issue_row(
                        source_file,
                        row_number,
                        "InvalidMidtermGrade",
                        f"Sheet '{worksheet.title}' has a blank or non-numeric MidtermGrade.",
                        raw_student_number,
                        raw_student_name,
                        raw_midterm_exam,
                        raw_midterm_grade,
                    )
                )
                continue

            notes = []
            if filename_note:
                notes.append(filename_note)

            student_number = find_masterlist_student_number(masterlist, student_name)
            if not is_blank(student_number):
                notes.append("StudentNumber filled from masterlist.")
            else:
                student_number = ""
                notes.append("StudentNumber not found in masterlist.")

            result_row = {
                "SourceFile": source_file,
                "Section": section,
                "Course": course,
                "StudentNumber": student_number,
                "StudentName": student_name,
                "NormalizedName": normalize_name(student_name),
                "MidtermExam": raw_midterm_exam,
                "MidtermGrade": grade,
                "Status": "Passed" if grade >= passing_grade else "Fail",
                "Notes": " ".join(notes),
                "RawRowJson": json.dumps(row_values, default=str),
            }
            result_row.update(standard_grade_values)
            result_row["MidtermGrade"] = grade
            result_rows.append(result_row)

    workbook.close()
    source_info["Processed"] = processed_any_sheet

    if not processed_any_sheet:
        issues.append(
            issue_row(
                source_file,
                "",
                "NoProcessableSheet",
                "No worksheet contained all required processable columns.",
            )
        )

    return result_rows, issues, source_info


def apply_duplicate_policy(
    result_rows: List[Dict[str, Any]], issues: List[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    kept_rows: List[Dict[str, Any]] = []
    row_indexes: Dict[Tuple[str, str], int] = {}

    for result in result_rows:
        normalized_name = normalize_name(result.get("StudentName"))
        course = str(result.get("Course") or "")
        key = (normalized_name, course)
        if normalized_name and course and key in row_indexes:
            previous_index = row_indexes[key]
            issues.append(
                issue_row(
                    result.get("SourceFile") or "",
                    "",
                    "DuplicateStudentCourse",
                    "Duplicate student/course found in one batch; keeping the last record.",
                    result.get("StudentNumber"),
                    result.get("StudentName"),
                    result.get("MidtermExam"),
                    result.get("MidtermGrade"),
                )
            )
            kept_rows[previous_index] = result
            continue

        row_indexes[key] = len(kept_rows)
        kept_rows.append(result)

    return kept_rows


def ingest_folder(
    input_folder: Path, masterlist_path: Path, passing_grade: float
) -> IngestResult:
    result = IngestResult(passing_grade=passing_grade)
    masterlist, masterlist_records, masterlist_issues = load_masterlist(masterlist_path)
    result.masterlist_records = masterlist_records
    result.issues.extend(masterlist_issues)

    if not input_folder.exists():
        result.issues.append(
            issue_row(
                "",
                "",
                "InputFolderMissing",
                f"Input folder does not exist: {input_folder}",
            )
        )
        return result

    workbook_paths = sorted(
        path
        for path in input_folder.iterdir()
        if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm"}
    )

    if not workbook_paths:
        result.issues.append(
            issue_row(
                "",
                "",
                "NoInputFiles",
                f"No .xlsx or .xlsm files found in: {input_folder}",
            )
        )
        return result

    for workbook_path in workbook_paths:
        workbook_results, workbook_issues, source_info = process_workbook(
            workbook_path, passing_grade, masterlist
        )
        if source_info.get("Processed"):
            result.files_processed += 1
        result.source_files.append(source_info)
        result.result_rows.extend(workbook_results)
        result.issues.extend(workbook_issues)

    result.result_rows = apply_duplicate_policy(result.result_rows, result.issues)
    return result
