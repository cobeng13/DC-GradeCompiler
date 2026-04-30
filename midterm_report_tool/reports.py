import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook

from ingest import ISSUE_COLUMNS, PASSED_COLUMNS, is_blank, normalize_name

MATRIX_NAME_MERGE_MIN_ORDERED_CHARS = 10
MATRIX_NAME_MERGE_MIN_SHORTER_RATIO = 0.92
MATRIX_NAME_MERGE_MIN_LONGER_RATIO = 0.82
NAME_SUFFIX_TOKENS = {"jr", "sr", "ii", "iii", "iv", "v"}


def write_sheet(
    workbook: Workbook, title: str, columns: List[str], rows: List[Dict[str, Any]]
) -> None:
    worksheet = workbook.create_sheet(title=title)
    worksheet.append(columns)

    for row in rows:
        worksheet.append([row.get(column, "") for column in columns])

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 12), 60
        )


def year_sheet_name(section: Any) -> str:
    if is_blank(section):
        return "Unassigned"

    match = re.match(r"\s*(\d+)", str(section))
    if not match:
        return "Unassigned"
    return f"Year{match.group(1)}"


def safe_sheet_title(title: str, used_titles: set) -> str:
    safe_title = re.sub(r"[\[\]:*?/\\]", "_", title)[:31] or "Sheet"
    candidate = safe_title
    suffix = 1

    while candidate in used_titles:
        suffix_text = f"_{suffix}"
        candidate = f"{safe_title[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1

    used_titles.add(candidate)
    return candidate


def compact_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize_name(value))


def leading_name_token(normalized_name: str) -> str:
    return normalized_name.split()[0] if normalized_name.split() else ""


def significant_name_tokens(value: Any) -> List[str]:
    return [
        token
        for token in normalize_name(value).split()
        if token not in NAME_SUFFIX_TOKENS
    ]


def token_initials(tokens: List[str]) -> str:
    return "".join(token[0] for token in tokens if token)


def name_token_similarity(left: str, right: str) -> float:
    if left == right:
        return 1.0
    longer_length = max(len(left), len(right))
    if longer_length == 0:
        return 0.0
    matches = ordered_character_match_count(left, right)
    return matches / longer_length


def ordered_character_match_count(left: str, right: str) -> int:
    if not left or not right:
        return 0

    previous = [0] * (len(right) + 1)
    for left_character in left:
        current = [0]
        for index, right_character in enumerate(right, start=1):
            if left_character == right_character:
                current.append(previous[index - 1] + 1)
            else:
                current.append(max(previous[index], current[-1]))
        previous = current
    return previous[-1]


def sections_are_merge_compatible(left: Any, right: Any) -> bool:
    if is_blank(left) or is_blank(right):
        return True
    return str(left).strip().lower() == str(right).strip().lower()


def names_are_ordered_merge_match(
    left_name: Any,
    right_name: Any,
    min_ordered_chars: int = MATRIX_NAME_MERGE_MIN_ORDERED_CHARS,
    min_shorter_ratio: float = MATRIX_NAME_MERGE_MIN_SHORTER_RATIO,
    min_longer_ratio: float = MATRIX_NAME_MERGE_MIN_LONGER_RATIO,
) -> bool:
    left_normalized = normalize_name(left_name)
    right_normalized = normalize_name(right_name)
    if not left_normalized or not right_normalized or left_normalized == right_normalized:
        return False

    if leading_name_token(left_normalized) != leading_name_token(right_normalized):
        return False

    left_tokens = significant_name_tokens(left_normalized)
    right_tokens = significant_name_tokens(right_normalized)
    if len(left_tokens) < 3 or len(right_tokens) < 3:
        return False

    if left_tokens[0] != right_tokens[0]:
        return False

    shared_nonfamily_tokens = set(left_tokens[1:]) & set(right_tokens[1:])
    if len(shared_nonfamily_tokens) < 2:
        return False

    common_count = min(len(left_tokens), len(right_tokens))
    different_token_indexes = [
        index
        for index in range(common_count)
        if left_tokens[index] != right_tokens[index]
        and name_token_similarity(left_tokens[index], right_tokens[index]) < 0.82
    ]
    if len(different_token_indexes) > 1:
        return False

    if token_initials(left_tokens) != token_initials(right_tokens):
        return False

    left_compact = compact_name(left_normalized)
    right_compact = compact_name(right_normalized)
    shorter_length = min(len(left_compact), len(right_compact))
    longer_length = max(len(left_compact), len(right_compact))
    if shorter_length == 0:
        return False

    ordered_matches = ordered_character_match_count(left_compact, right_compact)
    return (
        ordered_matches >= min_ordered_chars
        and ordered_matches / shorter_length >= min_shorter_ratio
        and ordered_matches / longer_length >= min_longer_ratio
    )


def find_mergeable_student_index(
    students: List[Dict[str, Any]], normalized_name: str, section: Any
) -> Optional[int]:
    for index, student in enumerate(students):
        if not sections_are_merge_compatible(student.get("Section"), section):
            continue
        if names_are_ordered_merge_match(student.get("NormalizedName"), normalized_name):
            return index
    return None


def append_student_alias(student: Dict[str, Any], normalized_name: str) -> None:
    aliases = student.setdefault("NormalizedAliases", [])
    if normalized_name and normalized_name not in aliases:
        aliases.append(normalized_name)


def upsert_matrix_student(
    students: List[Dict[str, Any]],
    indexes: Dict[str, List[int]],
    student_name: Any,
    normalized_name: str,
    section: Any,
    from_masterlist: bool = False,
) -> None:
    if normalized_name in indexes:
        for index in indexes[normalized_name]:
            student = students[index]
            if sections_are_merge_compatible(student.get("Section"), section):
                if is_blank(student.get("Section")) and not is_blank(section):
                    student["Section"] = section
                return

    merge_index = find_mergeable_student_index(students, normalized_name, section)
    if merge_index is not None:
        indexes.setdefault(normalized_name, []).append(merge_index)
        student = students[merge_index]
        append_student_alias(student, normalized_name)
        if is_blank(student.get("Section")) and not is_blank(section):
            student["Section"] = section
        if not student.get("FromMasterlist") and len(str(student_name or "")) > len(
            str(student.get("StudentName") or "")
        ):
            student["StudentName"] = student_name
            student["NormalizedName"] = normalized_name
        if from_masterlist:
            student["FromMasterlist"] = True
        return

    indexes.setdefault(normalized_name, []).append(len(students))
    students.append(
        {
            "StudentName": student_name,
            "Section": section,
            "NormalizedName": normalized_name,
            "NormalizedAliases": [normalized_name],
            "FromMasterlist": from_masterlist,
        }
    )


def build_matrix_students(
    masterlist_records: List[Dict[str, Any]], result_rows: List[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    students: List[Dict[str, Any]] = []
    indexes: Dict[str, List[int]] = {}

    for record in masterlist_records:
        normalized_name = record.get("NormalizedName") or normalize_name(
            record.get("StudentName")
        )
        if not normalized_name:
            continue
        upsert_matrix_student(
            students,
            indexes,
            record.get("StudentName"),
            normalized_name,
            record.get("Section"),
            from_masterlist=True,
        )

    for result in result_rows:
        normalized_name = normalize_name(result.get("StudentName"))
        if not normalized_name:
            continue

        upsert_matrix_student(
            students,
            indexes,
            result.get("StudentName"),
            normalized_name,
            result.get("Section"),
        )

    return students


def matrix_courses(result_rows: List[Dict[str, Any]]) -> List[str]:
    courses = []
    seen_courses = set()
    for result in result_rows:
        course = result.get("Course")
        if is_blank(course) or course in seen_courses:
            continue
        seen_courses.add(course)
        courses.append(course)
    return courses


def write_matrix_report(
    output_file: Path,
    masterlist_records: List[Dict[str, Any]],
    result_rows: List[Dict[str, Any]],
    result_lookup: Dict[Tuple[str, str], Any],
) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    courses = matrix_courses(result_rows)

    students = build_matrix_students(masterlist_records, result_rows)
    students_by_sheet: Dict[str, List[Dict[str, Any]]] = {}
    for student in students:
        sheet_name = year_sheet_name(student.get("Section"))
        students_by_sheet.setdefault(sheet_name, []).append(student)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_titles = set()
    sheet_order = sorted(
        students_by_sheet,
        key=lambda name: (name == "Unassigned", name),
    )
    if not sheet_order:
        sheet_order = ["Unassigned"]
        students_by_sheet["Unassigned"] = []

    for sheet_name in sheet_order:
        worksheet = workbook.create_sheet(safe_sheet_title(sheet_name, used_titles))
        columns = ["Name", "Section"] + courses
        worksheet.append(columns)

        for student in sorted(
            students_by_sheet[sheet_name],
            key=lambda item: (
                str(item.get("Section") or ""),
                str(item.get("StudentName") or ""),
            ),
        ):
            normalized_aliases = student.get("NormalizedAliases") or [
                student.get("NormalizedName")
            ]
            row = [student.get("StudentName"), student.get("Section")]
            for course in courses:
                value = ""
                for normalized_name in normalized_aliases:
                    value = result_lookup.get((normalized_name, course), "")
                    if not is_blank(value):
                        break
                row.append(value)
            worksheet.append(row)

        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(max_length + 2, 12), 50
            )

    workbook.save(output_file)


def generate_pass_fail_matrix_report(
    output_file: Path,
    masterlist_records: List[Dict[str, Any]],
    result_rows: List[Dict[str, Any]],
) -> None:
    result_lookup: Dict[Tuple[str, str], str] = {}
    for result in result_rows:
        normalized_name = normalize_name(result.get("StudentName"))
        course = result.get("Course")
        status = result.get("Status")
        if normalized_name and not is_blank(course) and status in {"Passed", "Fail"}:
            result_lookup[(normalized_name, course)] = status

    write_matrix_report(output_file, masterlist_records, result_rows, result_lookup)


def generate_midterm_grade_matrix_report(
    output_file: Path,
    masterlist_records: List[Dict[str, Any]],
    result_rows: List[Dict[str, Any]],
) -> None:
    result_lookup: Dict[Tuple[str, str], Any] = {}
    for result in result_rows:
        normalized_name = normalize_name(result.get("StudentName"))
        course = result.get("Course")
        grade = result.get("MidtermGrade")
        if normalized_name and not is_blank(course) and not is_blank(grade):
            result_lookup[(normalized_name, course)] = grade

    write_matrix_report(output_file, masterlist_records, result_rows, result_lookup)


def generate_workbook_reports(
    output_file: Path,
    pass_fail_output_file: Path,
    grade_output_file: Path,
    masterlist_records: List[Dict[str, Any]],
    result_rows: List[Dict[str, Any]],
    issues: List[Dict[str, Any]],
) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    passed_students = [
        result for result in result_rows if result.get("Status") == "Passed"
    ]

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    write_sheet(workbook, "PassedStudents", PASSED_COLUMNS, passed_students)
    write_sheet(workbook, "Issues", ISSUE_COLUMNS, issues)
    workbook.save(output_file)

    generate_pass_fail_matrix_report(
        pass_fail_output_file, masterlist_records, result_rows
    )
    generate_midterm_grade_matrix_report(
        grade_output_file, masterlist_records, result_rows
    )
