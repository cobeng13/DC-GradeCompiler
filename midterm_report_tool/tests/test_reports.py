import sys
import shutil
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from reports import (
    build_matrix_students,
    generate_midterm_grade_matrix_report,
    names_are_ordered_merge_match,
    ordered_character_match_count,
)


class ReportMatrixMergeTests(unittest.TestCase):
    @contextmanager
    def workspace_tempdir(self):
        temp_root = Path(__file__).resolve().parents[1] / ".test_tmp"
        temp_root.mkdir(exist_ok=True)
        temp_dir = temp_root / f"report_{uuid.uuid4().hex}"
        temp_dir.mkdir()
        try:
            yield temp_dir
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    def test_ordered_character_match_counts_subsequence_overlap(self):
        self.assertEqual(
            ordered_character_match_count(
                "andayaisiahximuelcarandang", "andayaisiahximuel"
            ),
            len("andayaisiahximuel"),
        )

    def test_ordered_name_merge_matches_close_typo_variants(self):
        self.assertTrue(
            names_are_ordered_merge_match(
                "Culis, Mizzy Marimontt Bejasa", "Culis, Mizzy Mariomntt Bejasa"
            )
        )

    def test_ordered_name_merge_rejects_shortened_or_different_students(self):
        self.assertFalse(
            names_are_ordered_merge_match(
                "Andaya, Isiah Ximuel Carandang", "Andaya, Isiah Ximuel"
            )
        )
        self.assertFalse(
            names_are_ordered_merge_match("Anog, Alyssa Naga?o", "Anog, Alyssa")
        )
        self.assertFalse(
            names_are_ordered_merge_match(
                "Concepcion, Angel Tom Rasay", "Concepcion, Tom"
            )
        )
        self.assertFalse(
            names_are_ordered_merge_match("Mau, Bronny Junior", "Mau, Bronny")
        )
        self.assertFalse(
            names_are_ordered_merge_match(
                "Catibog, Danielle Ysabel Lanip", "Catibog, Denise Alexa Lanip"
            )
        )
        self.assertFalse(
            names_are_ordered_merge_match(
                "Bautista, Elaiza Kathrina Calayan",
                "Bautista, Eloiza Katherine Calayan",
            )
        )

    def test_matrix_students_merge_compatible_name_variants(self):
        students = build_matrix_students(
            [
                {
                    "StudentName": "Culis, Mizzy Marimontt Bejasa",
                    "Section": "1A",
                    "NormalizedName": "culis mizzy marimontt bejasa",
                }
            ],
            [
                {
                    "StudentName": "Culis, Mizzy Mariomntt Bejasa",
                    "Section": "1A",
                    "Course": "CourseA",
                    "MidtermGrade": 90,
                },
                {
                    "StudentName": "Culis, Mizzy Marimontt Bejasa",
                    "Section": "2A",
                    "Course": "CourseB",
                    "MidtermGrade": 91,
                },
            ],
        )

        self.assertEqual(len(students), 1)
        merged_student = students[0]
        self.assertEqual(
            set(merged_student["NormalizedAliases"]),
            {"culis mizzy marimontt bejasa", "culis mizzy mariomntt bejasa"},
        )
        self.assertEqual(set(merged_student["Sections"]), {"1A", "2A"})

    def test_matrix_export_places_irregular_student_in_each_year_sheet(self):
        with self.workspace_tempdir() as temp_dir:
            output_file = Path(temp_dir) / "grades.xlsx"
            generate_midterm_grade_matrix_report(
                output_file,
                [
                    {
                        "StudentName": "Irregular, Student Example",
                        "Section": "3A",
                        "NormalizedName": "irregular student example",
                    }
                ],
                [
                    {
                        "StudentName": "Irregular, Student Example",
                        "SourceFile": "1A_SecondYearCourse.xlsx",
                        "Section": "2A",
                        "Course": "SecondYearCourse",
                        "MidtermGrade": 88,
                    },
                    {
                        "StudentName": "Irregular, Student Example",
                        "SourceFile": "2A_ThirdYearCourse.xlsx",
                        "Section": "3A",
                        "Course": "ThirdYearCourse",
                        "MidtermGrade": 91,
                    },
                ],
            )

            workbook = load_workbook(output_file)
            year1_rows = list(workbook["Year1"].iter_rows(values_only=True))
            year2_rows = list(workbook["Year2"].iter_rows(values_only=True))

        year1_header = year1_rows[0]
        year1_student = year1_rows[1]
        year2_header = year2_rows[0]
        year2_student = year2_rows[1]

        self.assertEqual(year1_student[0], "Irregular, Student Example")
        self.assertEqual(year1_student[1], "1A")
        self.assertEqual(
            year1_student[year1_header.index("SecondYearCourse")],
            88,
        )
        self.assertIsNone(year1_student[year1_header.index("ThirdYearCourse")])

        self.assertEqual(year2_student[0], "Irregular, Student Example")
        self.assertEqual(year2_student[1], "2A")
        self.assertIsNone(year2_student[year2_header.index("SecondYearCourse")])
        self.assertEqual(
            year2_student[year2_header.index("ThirdYearCourse")],
            91,
        )


if __name__ == "__main__":
    unittest.main()
